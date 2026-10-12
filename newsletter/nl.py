"""Heirstone newsletter tool: render issues, send tests, draft broadcasts in Resend, manage the list.

    python newsletter/nl.py render  2026-11                 HTML + text previews in newsletter/out/
    python newsletter/nl.py test    2026-11 --to me@x.com   send both editions to one inbox (subject marked TEST)
    python newsletter/nl.py draft   2026-11                 create the EN and FR broadcasts in Resend (not sent)
    python newsletter/nl.py contacts template <file.xlsx>   create the contacts workbook (headers + guide)
    python newsletter/nl.py contacts sync <file.xlsx|csv> [--dry-run]   add/update contacts from the workbook
    python newsletter/nl.py contacts export <file.xlsx>     download the live list (incl. website sign-ups)
    python newsletter/nl.py setup                           create the segment, topics and fields in Resend

Resend calls need a full-access API key in RESEND_API_KEY (or in ~/.heirstone/resend_api_key).
Sending is always a person's click in the Resend dashboard: this tool only creates drafts.
See newsletter/README.md for the whole design.
"""
import argparse
import csv
import html
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
CONFIG_FILE = HERE / 'config.json'
OUT = HERE / 'out'
CATALOG = json.loads((ROOT / 'private' / 'reports' / 'catalog.json').read_text(encoding='utf-8'))
EMAIL_RE = re.compile(r'^[^\s@<>"]+@[^\s@<>"]+\.[^\s@<>"]{2,}$')

# ---------------------------------------------------------------- brand
INK, NAVY, SLATE, MIST, TINT = '#0E1B2A', '#13273A', '#5D6E82', '#97A6B6', '#8FA0B2'
TEXT, TEXT2, LIGHT = '#263340', '#3C4956', '#6E7F92'
PAPER, CARD, LINE = '#F6F4EF', '#FCFBF8', '#E4DED2'
SERIF = "'Playfair Display',Georgia,'Times New Roman',serif"
SANS = "'Libre Franklin',Arial,Helvetica,sans-serif"
INNER = 518  # 600 container - 2 x 40 padding - 2 x 1 border

MONTHS = {
    'en': 'January February March April May June July August September October November December'.split(),
    'fr': 'janvier février mars avril mai juin juillet août septembre octobre novembre décembre'.split(),
}
# Fixed wording of the template. Issue files hold everything that changes month to month.
LABELS = {
    'en': {
        'tagline': 'Research on real estate, tourism and national vision across Africa, the Middle East and the Mediterranean',
        'issue': 'No.', 'dear': 'Dear', 'reader': 'reader',
        'lead': 'This month’s report', 'brief': 'In brief', 'more': 'More from Heirstone Research', 'read': 'Read the summary',
        'source': 'Source', 'page': 'p.', 'series': {'africa': 'Africa Series', None: 'Middle East & Mediterranean'},
        'cover': 'Cover of the report',
        'links': [('Reports', '/pages/reports'), ('Insights', '/pages/insights'), ('Contact', '/pages/contact')],
        'linkedin': 'LinkedIn',
        'why': 'You are receiving {name} because you requested a Heirstone report, asked to receive our research or work with us.',
        'unsub': 'Unsubscribe or manage your preferences',
        'switch': 'Prefer the French edition? Reply to this email with “French”.',
        'offices': 'Dubai · Cairo',
    },
    'fr': {
        'tagline': 'Recherche en immobilier, tourisme et vision nationale en Afrique, au Moyen-Orient et en Méditerranée',
        'issue': 'N°', 'dear': 'Bonjour', 'reader': 'à toutes et à tous',
        'lead': 'Le rapport du mois', 'brief': 'En bref', 'more': 'Également chez Heirstone Research', 'read': 'Lire la présentation',
        'source': 'Source', 'page': 'p.', 'series': {'africa': 'Série Afrique', None: 'Moyen-Orient et Méditerranée'},
        'cover': 'Couverture du rapport',
        'links': [('Rapports', '/fr/pages/reports'), ('Analyses', '/fr/pages/insights'), ('Contact', '/fr/pages/contact')],
        'linkedin': 'LinkedIn',
        'why': 'Vous recevez {name} parce que vous avez demandé un rapport Heirstone, souhaité recevoir nos publications ou travaillez avec nous.',
        'unsub': 'Se désinscrire ou gérer vos préférences',
        'switch': 'Vous préférez l’édition anglaise ? Répondez « anglais » à cet e-mail.',
        'offices': 'Dubaï · Le Caire',
    },
}
LINKEDIN = 'https://www.linkedin.com/company/heirstoneconsulting'


def load_config():
    return json.loads(CONFIG_FILE.read_text(encoding='utf-8'))


# ---------------------------------------------------------------- report data (catalog + landing pages)
class _Meta(HTMLParser):
    def __init__(self):
        super().__init__()
        self.description = ''

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'meta' and a.get('name') == 'description' and not self.description:
            self.description = a.get('content') or ''


def report_info(slug, lang):
    """Title, series, landing page and summary of a report in one language, from the catalog and its page."""
    if slug not in CATALOG:
        sys.exit(f'Unknown report "{slug}". Known: {", ".join(sorted(CATALOG))}')
    r = CATALOG[slug]
    ed = (r.get('editions') or {}).get(lang) or (r.get('editions') or {}).get('en') or r
    page = f"{'/' + lang if lang != 'en' else ''}/pages/report-{slug}"
    meta = _Meta()
    page_file = ROOT / (page.lstrip('/') + '.html')
    if page_file.exists():
        meta.feed(page_file.read_text(encoding='utf-8'))
    has_lang = bool((r.get('editions') or {}).get(lang))
    return {
        'slug': slug, 'series': r.get('series'), 'sector': ed.get('sector') or r['sector'], 'title': ed.get('title') or r['title'],
        'page': page, 'summary': meta.description, 'pages': ed.get('pages'),
        'cover': f"/Media/email/covers/{slug}{'.fr' if lang == 'fr' and has_lang else ''}.jpg",
    }


# ---------------------------------------------------------------- email images
def ensure_images(slugs, langs):
    """Email-safe JPG covers (Outlook does not show WebP). EN from the site cover, FR from the French PDF."""
    from PIL import Image
    made = []
    covers = ROOT / 'Media' / 'email' / 'covers'
    covers.mkdir(parents=True, exist_ok=True)
    for slug in slugs:
        target = covers / f'{slug}.jpg'
        if not target.exists():
            im = Image.open(ROOT / 'Media' / 'reports' / f'{slug}.webp').convert('RGB')
            im.resize((360, round(360 * im.height / im.width)), Image.LANCZOS).save(target, 'JPEG', quality=82, optimize=True, progressive=True)
            made.append(target)
        fr_pdf = ROOT / 'private' / 'reports' / f'{slug}.fr.pdf'
        fr_target = covers / f'{slug}.fr.jpg'
        if 'fr' in langs and fr_pdf.exists() and not fr_target.exists():
            import fitz  # PyMuPDF
            with fitz.open(fr_pdf) as doc:
                page = doc[0]
                pix = page.get_pixmap(matrix=fitz.Matrix(720 / page.rect.width, 720 / page.rect.width))
                im = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            im.resize((360, round(360 * im.height / im.width)), Image.LANCZOS).save(fr_target, 'JPEG', quality=82, optimize=True, progressive=True)
            made.append(fr_target)
    return made


# ---------------------------------------------------------------- HTML helpers
def esc(s):
    return html.escape(str(s), quote=True)


def md(s, link_style=''):
    """Escape, then allow **bold** and [text](url) in issue text."""
    out = esc(s)
    out = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', out)
    return re.sub(r'\[(.+?)\]\((.+?)\)', lambda m: f'<a href="{m.group(2)}" style="{link_style}">{m.group(1)}</a>', out)


def plain(s):
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', str(s))
    return re.sub(r'\[(.+?)\]\((.+?)\)', r'\1 (\2)', s)


def track(url, site, campaign, content=''):
    """Absolute URL with newsletter UTM tags (site links only)."""
    if url.startswith('/'):
        url = site + url
    if not url.startswith(site):
        return url
    parts = urllib.parse.urlsplit(url)
    q = urllib.parse.parse_qsl(parts.query)
    q += [('utm_source', 'newsletter'), ('utm_medium', 'email'), ('utm_campaign', campaign)] + ([('utm_content', content)] if content else [])
    return urllib.parse.urlunsplit(parts._replace(query=urllib.parse.urlencode(q)))


def p(text, size=16, line=26, color=TEXT, margin='0 0 16px 0', family=SANS, extra=''):
    cls = 'serif' if family == SERIF else 'sans'
    return f'<p class="{cls}" style="margin:{margin};font-family:{family};font-size:{size}px;line-height:{line}px;color:{color};{extra}">{text}</p>'


def label(text, color=SLATE, margin='0 0 18px 0'):
    return p(esc(text), 11, 16, color, margin, SANS, 'letter-spacing:3px;text-transform:uppercase;font-weight:600;')


def button(text, url, bg=INK, fg='#FFFFFF'):
    return (f'<table role="presentation" cellpadding="0" cellspacing="0" border="0"><tr>'
            f'<td bgcolor="{bg}" style="background-color:{bg};padding:13px 24px;">'
            f'<a href="{esc(url)}" class="sans" style="font-family:{SANS};font-size:14px;line-height:18px;font-weight:600;color:{fg};text-decoration:none;display:inline-block;">{esc(text)}</a>'
            f'</td></tr></table>')


def row(inner, pad='32px 40px 0 40px', bg=None):
    bga = f' bgcolor="{bg}"' if bg else ''
    bgs = f'background-color:{bg};' if bg else ''
    return f'<tr><td class="px"{bga} style="{bgs}padding:{pad};">{inner}</td></tr>'


def rule():
    return row(f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0"><tr><td style="border-top:1px solid {LINE};font-size:0;line-height:0;height:1px;">&nbsp;</td></tr></table>', '32px 40px 0 40px')


def columns(cells):
    """Fluid-hybrid columns: side by side in a 600px window, stacked on phones; real tables in Outlook."""
    mso_open = '<!--[if mso]><table role="presentation" width="%d" cellpadding="0" cellspacing="0" border="0"><tr><![endif]-->' % INNER
    out = [mso_open]
    for width, inner in cells:
        out.append(f'<!--[if mso]><td width="{width}" valign="top"><![endif]-->'
                   f'<div class="stack" style="display:inline-block;width:100%;max-width:{width}px;vertical-align:top;">{inner}</div>'
                   '<!--[if mso]></td><![endif]-->')
    out.append('<!--[if mso]></tr></table><![endif]-->')
    return f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0"><tr><td valign="top" style="font-size:0;line-height:0;">{"".join(out)}</td></tr></table>'


# ---------------------------------------------------------------- render one edition
def render(issue, lang, cfg):
    e = issue['editions'][lang]
    L = LABELS[lang]
    site = cfg['site'].rstrip('/')
    name = cfg['name'][lang]
    camp = issue.get('campaign') or f"brief-{issue['date'][:7]}"
    d = date.fromisoformat(issue['date'])
    when = f"{MONTHS[lang][d.month - 1]} {d.year}"
    number = f"{L['issue']} {int(issue['number']):02d}"
    link = f'color:{INK};text-decoration:underline;'
    T = []  # plain-text version, built alongside
    rows = []

    # Masthead
    mast = (
        '<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0"><tr>'
        f'<td width="44" valign="middle" style="width:44px;"><a href="{esc(track(site + ("/fr/" if lang == "fr" else "/"), site, camp, "masthead"))}">'
        f'<img src="{site}/Media/email/heirstone-mark.png" width="44" height="44" alt="Heirstone Consulting" border="0" style="display:block;width:44px;height:44px;border:0;"></a></td>'
        f'<td valign="middle" class="sans" style="padding-left:14px;font-family:{SANS};font-size:11px;line-height:16px;letter-spacing:3px;color:{TINT};text-transform:uppercase;">Heirstone Consulting<br>Research</td>'
        '</tr></table>'
        f'<h1 class="serif h1" style="margin:30px 0 10px 0;font-family:{SERIF};font-size:40px;line-height:46px;font-weight:400;color:#FFFFFF;">{esc(name)}</h1>'
        + p(f'{esc(number)} &nbsp;·&nbsp; {esc(when)}', 11, 16, TINT, '0 0 12px 0', SANS, 'letter-spacing:2px;text-transform:uppercase;')
        + p(esc(L['tagline']), 14, 21, MIST, '0')
    )
    rows.append(row(mast, '28px 40px 32px 40px', INK))
    T += [f"{name.upper()}  ·  {number}  ·  {when}", L['tagline'], '']

    # Greeting, editor's note, sign-off
    greet = f"{L['dear']} {{{{{{contact.first_name|{L['reader']}}}}}}},"
    body = p(esc(greet))
    T += [greet, '']
    for para in e.get('intro', []):
        body += p(md(para, link))
        T += [plain(para), '']
    if e.get('signoff'):
        s = e['signoff']
        body += p(f"<strong>{esc(s['name'])}</strong><br><span style=\"color:{LIGHT};\">{esc(s.get('role', ''))}</span>", 15, 22, TEXT, '4px 0 0 0')
        T += [s['name'], s.get('role', ''), '']
    rows.append(row(body, '36px 40px 4px 40px'))

    # Lead report
    if e.get('lead'):
        lead = e['lead']
        r = report_info(lead['report'], lang)
        get = track(f"{r['page']}?report={r['slug']}", site, camp, 'lead-cover')
        kicker = f"{L['series'][r['series']]} · {r['sector']}"
        cover = (f'<div style="padding:0 0 22px 0;"><a href="{esc(get)}"><img src="{site}{r["cover"]}" width="160" height="{round(160 * 1274 / 900)}" '
                 f'alt="{esc(L["cover"])} {esc(r["title"])}" border="0" style="display:block;width:160px;height:auto;border:1px solid {LINE};"></a></div>')
        text = (p(esc(kicker), 11, 17, SLATE, '0 0 10px 0', SANS, 'letter-spacing:2px;text-transform:uppercase;font-weight:600;')
                + f'<h2 class="serif" style="margin:0 0 14px 0;font-family:{SERIF};font-size:28px;line-height:34px;font-weight:400;color:{INK};">{esc(r["title"])}</h2>'
                + p(md(lead['dek'], link), 15, 24, TEXT2, '0 0 22px 0')
                + button(lead.get('cta') or 'Get the report', track(f"{r['page']}?report={r['slug']}", site, camp, 'lead-button')))
        rows.append(row(label(L['lead']) + columns([(180, cover), (338, text)])))
        T += ['-' * 60, L['lead'].upper(), kicker, r['title'], '', plain(lead['dek']), '',
              f"{lead.get('cta') or 'Get the report'}: {track(r['page'] + '?report=' + r['slug'], site, camp, 'lead-button')}", '']

        stats = lead.get('stats', [])
        if stats:
            w = INNER // len(stats)
            cells = []
            for i, s in enumerate(stats):
                src = f"{L['page']} {s['page']}" if s.get('page') else ''
                cells.append((w, (
                    f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0"><tr>'
                    f'<td class="stat" style="border-top:2px solid {INK};padding:14px {"14px" if i < len(stats) - 1 else "0"} 0 0;">'
                    f'<p class="serif" style="margin:0 0 6px 0;font-family:{SERIF};font-size:30px;line-height:34px;color:{INK};">{esc(s["value"])}</p>'
                    + p(md(s['label'], link), 13, 19, TEXT2, '0 0 4px 0')
                    + (p(esc(src), 11, 16, LIGHT, '0') if src else '')
                    + '</td></tr></table>')))
            rows.append(row(columns(cells), '30px 40px 0 40px'))
            T += [f"  {s['value']}  {plain(s['label'])}" + (f" ({L['page']} {s['page']})" if s.get('page') else '') for s in stats]
            T.append('')

    # In brief
    if e.get('brief'):
        items = ''
        T += ['-' * 60, L['brief'].upper(), '']
        for i, b in enumerate(e['brief'], 1):
            r = report_info(b['report'], lang) if b.get('report') else None
            src = ''
            if r:
                where = f", {L['page']} {b['page']}" if b.get('page') else ''
                colon = ' : ' if lang == 'fr' else ': '
                ref = f'<a href="{esc(track(r["page"], site, camp, f"brief-{i}"))}" style="color:{LIGHT};text-decoration:underline;">{esc(r["title"])}</a>'
                src = p(f"{esc(L['source'])}{colon}{ref}{esc(where)}", 12, 18, LIGHT, '0')
            items += (f'<tr><td width="46" valign="top" class="serif" style="width:46px;padding:0 0 22px 0;font-family:{SERIF};font-size:22px;line-height:26px;color:{TINT};">{i:02d}</td>'
                      f'<td valign="top" style="padding:0 0 22px 0;">'
                      + p(esc(b['title']), 17, 24, INK, '0 0 6px 0', SERIF)
                      + p(md(b['text'], link), 15, 24, TEXT2, '0 0 6px 0') + src + '</td></tr>')
            T += [f"{i:02d}  {b['title']}", plain(b['text'])] + ([f"     {L['source']}: {r['title']}" + (f", {L['page']} {b['page']}" if b.get('page') else '')] if r else []) + ['']
        rows.append(rule())
        rows.append(row(label(L['brief']) + f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">{items}</table>', '28px 40px 0 40px'))

    # More reports
    if e.get('more'):
        items = ''
        T += ['-' * 60, L['more'].upper(), '']
        for i, slug in enumerate(e['more']):
            r = report_info(slug, lang)
            url = track(r['page'], site, camp, f'more-{slug}')
            top = f'border-top:1px solid {LINE};' if i else ''
            items += (f'<tr><td width="58" valign="top" style="width:58px;padding:16px 0 16px 0;{top}">'
                      f'<a href="{esc(url)}"><img src="{site}{r["cover"]}" width="44" height="62" alt="{esc(L["cover"])} {esc(r["title"])}" border="0" style="display:block;width:44px;height:62px;border:1px solid {LINE};"></a></td>'
                      f'<td valign="top" style="padding:16px 0 16px 0;{top}">'
                      + p(esc(r['sector']), 11, 16, SLATE, '0 0 4px 0', SANS, 'letter-spacing:2px;text-transform:uppercase;font-weight:600;')
                      + p(f'<a href="{esc(url)}" style="color:{INK};text-decoration:none;">{esc(r["title"])}</a>', 19, 25, INK, '0 0 6px 0', SERIF)
                      + (p(esc(r['summary']), 13, 20, LIGHT, '0') if r['summary'] else '') + '</td></tr>')
            T += [f"{r['sector']}: {r['title']}", r['summary'], url, '']
        rows.append(rule())
        rows.append(row(label(L['more'], margin='0 0 4px 0') + f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0">{items}</table>', '28px 40px 0 40px'))

    # Note
    if e.get('note'):
        n = e['note']
        box = (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0"><tr><td bgcolor="{PAPER}" style="background-color:{PAPER};padding:24px 26px;">'
               + p(esc(n['title']), 20, 26, INK, '0 0 10px 0', SERIF)
               + ''.join(p(md(t, link), 15, 24, TEXT2, '0 0 10px 0' if j < len(n['text']) - 1 else '0') for j, t in enumerate(n['text']))
               + '</td></tr></table>')
        rows.append(row(box, '28px 40px 0 40px'))
        T += ['-' * 60, n['title'].upper(), ''] + [plain(t) + '\n' for t in n['text']]

    # Closing call to action
    if e.get('closing'):
        c = e['closing']
        url = track(c['url'], site, camp, 'closing')
        inner = (p(esc(c['title']), 24, 30, '#FFFFFF', '0 0 12px 0', SERIF)
                 + p(md(c['text'], f'color:#FFFFFF;text-decoration:underline;'), 15, 24, '#C5CED8', '0 0 22px 0')
                 + button(c['cta'], url, PAPER, INK))
        rows.append(row('', '40px 40px 0 40px'))
        rows.append(row(inner, '36px 40px 38px 40px', NAVY))
        T += ['-' * 60, c['title'], plain(c['text']), f"{c['cta']}: {url}", '']

    # Footer
    links = ' &nbsp;·&nbsp; '.join(f'<a href="{esc(track(u, site, camp, "footer"))}" style="color:{SLATE};text-decoration:underline;">{esc(t)}</a>' for t, u in L['links'])
    links += f' &nbsp;·&nbsp; <a href="{LINKEDIN}" style="color:{SLATE};text-decoration:underline;">{L["linkedin"]}</a>'
    foot = (p(f'<strong style="color:{TEXT};">Heirstone Consulting</strong> &nbsp;·&nbsp; {esc(L["offices"])} &nbsp;·&nbsp; <a href="{site}" style="color:{SLATE};text-decoration:none;">heirstoneconsulting.com</a>', 12, 19, LIGHT, '0 0 8px 0')
            + p(links, 12, 19, LIGHT, '0 0 16px 0')
            + p(esc(L['why'].format(name=name)), 12, 19, LIGHT, '0 0 8px 0')
            + p(esc(L['switch']), 12, 19, LIGHT, '0 0 8px 0')
            + p(f'<a href="{{{{{{RESEND_UNSUBSCRIBE_URL}}}}}}" style="color:{SLATE};text-decoration:underline;">{esc(L["unsub"])}</a>', 12, 19, LIGHT, '0')
            + (p(esc(cfg['postal_address']), 12, 19, LIGHT, '8px 0 0 0') if cfg.get('postal_address') else ''))
    rows.append(row(foot, '30px 40px 34px 40px'))
    T += ['-' * 60, f"Heirstone Consulting · {L['offices']} · {site}", L['why'].format(name=name), L['switch'],
          f"{L['unsub']}: {{{{{{RESEND_UNSUBSCRIBE_URL}}}}}}"] + ([cfg['postal_address']] if cfg.get('postal_address') else [])

    pre = esc(e.get('preheader', ''))
    doc = f'''<!DOCTYPE html>
<html lang="{lang}" xmlns="http://www.w3.org/1999/xhtml" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta http-equiv="X-UA-Compatible" content="IE=edge">
<meta name="x-apple-disable-message-reformatting">
<meta name="format-detection" content="telephone=no, date=no, address=no, email=no">
<meta name="color-scheme" content="light">
<meta name="supported-color-schemes" content="light">
<title>{esc(e['subject'])}</title>
<!--[if !mso]><!--><link href="https://fonts.googleapis.com/css2?family=Libre+Franklin:wght@400;600&amp;family=Playfair+Display:wght@400&amp;display=swap" rel="stylesheet"><!--<![endif]-->
<!--[if mso]><noscript><xml><o:OfficeDocumentSettings><o:PixelsPerInch>96</o:PixelsPerInch></o:OfficeDocumentSettings></xml></noscript>
<style>.serif{{font-family:Georgia,serif !important;}} .sans{{font-family:Arial,sans-serif !important;}}</style><![endif]-->
<style>
:root{{color-scheme:light;supported-color-schemes:light;}}
body{{margin:0;padding:0;-webkit-text-size-adjust:100%;-ms-text-size-adjust:100%;}}
@media only screen and (max-width:620px){{
  .px{{padding-left:22px !important;padding-right:22px !important;}}
  .h1{{font-size:32px !important;line-height:38px !important;}}
  .stack{{display:block !important;width:100% !important;max-width:100% !important;}}
  .stat{{padding-right:0 !important;padding-bottom:16px !important;}}
}}
</style>
</head>
<body style="margin:0;padding:0;background-color:{PAPER};">
<div style="display:none;max-height:0;overflow:hidden;mso-hide:all;font-size:1px;line-height:1px;color:{PAPER};opacity:0;">{pre}{'&#847;&zwnj;&nbsp;' * 60}</div>
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" bgcolor="{PAPER}" style="background-color:{PAPER};">
<tr><td align="center" style="padding:24px 8px 24px 8px;">
<!--[if mso]><table role="presentation" width="600" cellpadding="0" cellspacing="0" border="0"><tr><td><![endif]-->
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="max-width:600px;background-color:{CARD};border:1px solid {LINE};">
{chr(10).join(rows)}
</table>
<!--[if mso]></td></tr></table><![endif]-->
</td></tr>
</table>
</body>
</html>
'''
    return doc, '\n'.join(T).strip() + '\n'


def load_issue(ref):
    path = Path(ref)
    if not path.exists():
        path = HERE / 'issues' / (ref if ref.endswith('.json') else ref + '.json')
    if not path.exists():
        sys.exit(f'Issue file not found: {ref}')
    issue = json.loads(path.read_text(encoding='utf-8'))
    issue['_id'] = path.stem
    return issue


def issue_reports(issue):
    slugs = set()
    for e in issue['editions'].values():
        if e.get('lead'):
            slugs.add(e['lead']['report'])
        slugs.update(b['report'] for b in e.get('brief', []) if b.get('report'))
        slugs.update(e.get('more', []))
    return sorted(slugs)


def build(issue, cfg):
    made = ensure_images(issue_reports(issue), issue['editions'].keys())
    OUT.mkdir(exist_ok=True)
    out = {}
    for lang in issue['editions']:
        doc, text = render(issue, lang, cfg)
        (OUT / f"{issue['_id']}.{lang}.html").write_text(doc, encoding='utf-8')
        (OUT / f"{issue['_id']}.{lang}.txt").write_text(text, encoding='utf-8')
        # Preview copy for a browser: images from this checkout (they may not be deployed yet), greeting fallback
        local = personalise(doc, '', '#unsubscribe').replace(cfg['site'].rstrip('/') + '/Media/', '../../Media/')
        (OUT / f"{issue['_id']}.{lang}.preview.html").write_text(local, encoding='utf-8')
        out[lang] = (doc, text)
    return out, made


def personalise(s, first_name, unsub_url):
    s = re.sub(r'\{\{\{contact\.first_name\|([^}]*)\}\}\}', lambda m: first_name or m.group(1), s)
    return s.replace('{{{RESEND_UNSUBSCRIBE_URL}}}', unsub_url)


# ---------------------------------------------------------------- Resend API
def api_key():
    key = os.environ.get('RESEND_API_KEY', '').strip()
    f = Path.home() / '.heirstone' / 'resend_api_key'
    if not key and f.exists():
        key = f.read_text(encoding='utf-8').strip()
    if not key:
        sys.exit('No Resend API key. Set RESEND_API_KEY (full access) or save it in ~/.heirstone/resend_api_key')
    return key


def resend(method, path, body=None):
    base = os.environ.get('RESEND_API_URL', 'https://api.resend.com').rstrip('/')
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(base + path, data=data, method=method, headers={
        'Authorization': 'Bearer ' + api_key(), 'Content-Type': 'application/json', 'User-Agent': 'heirstone-newsletter/1.0'})
    for attempt in range(4):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                time.sleep(0.12)  # stay under 10 requests a second
                raw = r.read().decode()
                return r.status, json.loads(raw) if raw else {}
        except urllib.error.HTTPError as e:
            raw = e.read().decode(errors='replace')
            if e.code == 429 and attempt < 3:
                time.sleep(1.5 * (attempt + 1))
                continue
            try:
                return e.code, json.loads(raw)
            except ValueError:
                return e.code, {'message': raw[:300]}
    return 429, {'message': 'rate limited'}


def need_ids(cfg):
    ids = cfg.get('resend') or {}
    if not ids.get('segment_id'):
        sys.exit('The newsletter segment is not set up yet: run `python newsletter/nl.py setup` (or fill newsletter/config.json).')
    return ids


# ---------------------------------------------------------------- commands: issues
def cmd_render(a):
    cfg = load_config()
    issue = load_issue(a.issue)
    out, made = build(issue, cfg)
    for lang in out:
        print(f"{lang}: {OUT / (issue['_id'] + '.' + lang + '.html')}  ({len(out[lang][0]) // 1024} KB)")
    if made:
        print(f'{len(made)} new email image(s) in Media/email/covers/; they must be deployed before the issue is sent.')


def cmd_test(a):
    cfg = load_config()
    issue = load_issue(a.issue)
    out, _ = build(issue, cfg)
    for lang in (a.lang.split(',') if a.lang else out):
        doc, text = out[lang]
        e = issue['editions'][lang]
        unsub = cfg['site'] + '/#test-unsubscribe-link'
        status, res = resend('POST', '/emails', {
            'from': cfg['from'], 'to': [a.to], 'reply_to': cfg['reply_to'], 'subject': f"[TEST {lang.upper()}] {e['subject']}",
            'html': personalise(doc, a.name, unsub), 'text': personalise(text, a.name, unsub)})
        print(f'{lang}: {"sent" if status < 300 else "FAILED"} ({status}) {res.get("id") or res.get("message", "")}')


def cmd_draft(a):
    cfg = load_config()
    ids = need_ids(cfg)
    issue = load_issue(a.issue)
    record = OUT / f"{issue['_id']}.broadcasts.json"
    if record.exists() and not a.force:
        sys.exit(f'Drafts already exist for this issue ({record.name}). Edit them in Resend, or pass --force to create new ones.')
    out, made = build(issue, cfg)
    # Images must already be live: the email loads them from the website
    missing = []
    for path in sorted({m for doc, _ in out.values() for m in re.findall(r'src="(https://[^"]+)"', doc)}):
        try:
            with urllib.request.urlopen(urllib.request.Request(path, method='HEAD', headers={'User-Agent': 'heirstone-newsletter/1.0'}), timeout=15) as r:
                if r.status != 200:
                    missing.append(path)
        except Exception:
            missing.append(path)
    if missing and not a.skip_image_check:
        sys.exit('These images are not live on the website yet. Deploy them first:\n  ' + '\n  '.join(missing))
    created = {}
    for lang, (doc, text) in out.items():
        topic = (ids.get('topics') or {}).get(lang)
        if not topic:
            print(f'{lang}: skipped (no topic id for this language in config.json)')
            continue
        e = issue['editions'][lang]
        status, res = resend('POST', '/broadcasts', {
            'segment_id': ids['segment_id'], 'topic_id': topic, 'from': cfg['from'], 'reply_to': cfg['reply_to'],
            'subject': e['subject'], 'html': doc, 'text': text,
            'name': f"{cfg['name'][lang]} {int(issue['number']):02d} ({lang.upper()}, {issue['date'][:7]})"})
        if status >= 300:
            sys.exit(f'{lang}: Resend refused the broadcast ({status}): {res}')
        created[lang] = res.get('id')
        print(f'{lang}: draft created, id {res.get("id")}')
    record.write_text(json.dumps({'created': datetime.now(timezone.utc).isoformat(), 'broadcasts': created}, indent=2), encoding='utf-8')
    print('\nNothing has been sent. Open https://resend.com/broadcasts, check each draft, send a test, then Send or Schedule.')


# ---------------------------------------------------------------- commands: contacts
COLUMNS = ['email', 'first_name', 'last_name', 'organisation', 'language', 'source', 'action', 'notes']
SOURCES = {
    'client': 'Current or past client',
    'met-in-person': 'Met in person and agreed to receive our research (meeting, event, card exchange)',
    'asked-to-subscribe': 'Asked to receive the newsletter (email, LinkedIn message, call)',
    'report-request': 'Requested a report (the website adds these automatically)',
}


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb['Contacts'] if 'Contacts' in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h or '').strip().lower().replace(' ', '_') for h in next(it)]
        return [dict(zip(head, ['' if v is None else str(v).strip() for v in r])) for r in it if any(v not in (None, '') for v in r)]
    with open(path, encoding='utf-8-sig', newline='') as f:
        return [{k.strip().lower().replace(' ', '_'): (v or '').strip() for k, v in r.items() if k} for r in csv.DictReader(f)]


_mx = {}


def has_mail_server(domain):
    """Does the domain accept mail? (MX lookup via Google public DNS; unknown counts as yes)."""
    if domain not in _mx:
        try:
            with urllib.request.urlopen(f'https://dns.google/resolve?name={urllib.parse.quote(domain)}&type=MX', timeout=10) as r:
                d = json.loads(r.read().decode())
            _mx[domain] = d.get('Status') == 0 and bool(d.get('Answer'))
            if not _mx[domain]:  # no MX: a plain A record can still receive mail
                with urllib.request.urlopen(f'https://dns.google/resolve?name={urllib.parse.quote(domain)}&type=A', timeout=10) as r:
                    _mx[domain] = bool(json.loads(r.read().decode()).get('Answer'))
        except Exception:
            _mx[domain] = True
    return _mx[domain]


def cmd_contacts_sync(a):
    cfg = load_config()
    ids = need_ids(cfg)
    topics = ids.get('topics') or {}
    rows = read_rows(a.file)
    now = datetime.now(timezone.utc).isoformat(timespec='seconds')
    log, seen = [], set()
    counts = {}

    def done(r, result, detail=''):
        counts[result] = counts.get(result, 0) + 1
        log.append({'email': r.get('email', ''), 'result': result, 'detail': detail})

    for r in rows:
        email = r.get('email', '').strip().lower()
        r['email'] = email
        action = r.get('action', '').strip().lower()
        lang = (r.get('language') or 'en').strip().lower()[:2]
        lang = lang if lang in ('en', 'fr') else 'en'
        source = r.get('source', '').strip().lower()
        if not EMAIL_RE.match(email):
            done(r, 'skipped', 'not a valid email address'); continue
        if email in seen:
            done(r, 'skipped', 'duplicate row in this file'); continue
        seen.add(email)
        if action not in ('', 'unsubscribe'):
            done(r, 'skipped', f'unknown action "{action}" (leave blank, or write unsubscribe)'); continue
        if action != 'unsubscribe' and source not in SOURCES:
            done(r, 'skipped', 'source missing or not one of: ' + ', '.join(SOURCES)); continue
        if action != 'unsubscribe' and not has_mail_server(email.split('@')[1]):
            done(r, 'skipped', 'the domain has no mail server'); continue
        if a.dry_run:
            done(r, 'ok (dry run)'); continue

        who = urllib.parse.quote(email)
        status, c = resend('GET', f'/contacts/{who}')
        if status == 200:
            if action == 'unsubscribe':
                s, res = resend('PATCH', f'/contacts/{who}', {'unsubscribed': True})
                done(r, 'unsubscribed' if s < 300 else 'error', '' if s < 300 else str(res)); continue
            if c.get('unsubscribed'):
                done(r, 'kept unsubscribed', 'this person unsubscribed; they are not re-added'); continue
            upd = {k: r[k] for k in ('first_name', 'last_name') if r.get(k)}
            if r.get('organisation'):
                upd['properties'] = {'organisation': r['organisation']}
            if upd:
                resend('PATCH', f'/contacts/{who}', upd)
            s, res = resend('POST', f'/contacts/{who}/segments/{ids["segment_id"]}')
            done(r, 'updated', 'already in Resend; name/organisation refreshed, edition preference kept'); continue
        if status != 404:
            done(r, 'error', f'lookup failed ({status}) {c}'); continue
        body = {
            'email': email, 'first_name': r.get('first_name', ''), 'last_name': r.get('last_name', ''),
            'unsubscribed': action == 'unsubscribe',
            'segments': [{'id': ids['segment_id']}],
            'topics': [{'id': topics[k], 'subscription': 'opt_in' if k == lang else 'opt_out'} for k in ('en', 'fr') if topics.get(k)],
            'properties': {k: v for k, v in {'organisation': r.get('organisation', ''), 'language': lang, 'source': source or 'import',
                                             'signed_up_at': now}.items() if v},
        }
        s, res = resend('POST', '/contacts', body)
        if s >= 400 and 'propert' in json.dumps(res).lower():
            body.pop('properties')
            s, res = resend('POST', '/contacts', body)
        if s < 300:
            done(r, 'unsubscribed' if action == 'unsubscribe' else 'added', f'{lang.upper()} edition' if action != 'unsubscribe' else 'kept on file so they are never re-added')
        else:
            done(r, 'error', str(res)[:200])

    out = Path(a.file).with_name(Path(a.file).stem + '.last-sync.csv')
    with open(out, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, ['email', 'result', 'detail'])
        w.writeheader()
        w.writerows(log)
    print(', '.join(f'{k}: {v}' for k, v in sorted(counts.items())) or 'No rows found.')
    print(f'Row-by-row result: {out}')


def cmd_contacts_export(a):
    cfg = load_config()
    ids = need_ids(cfg)
    rows, after = [], None
    while True:
        q = {'segment_id': ids['segment_id'], 'limit': 100}
        if after:
            q['after'] = after
        status, res = resend('GET', '/contacts?' + urllib.parse.urlencode(q))
        if status >= 300:
            sys.exit(f'Resend refused the list ({status}): {res}')
        data = res.get('data', [])
        rows += data
        if not res.get('has_more') or not data:
            break
        after = data[-1]['id']
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Live list'
    head = ['email', 'first_name', 'last_name', 'unsubscribed', 'created_at']
    ws.append(head)
    for c in rows:
        ws.append([c.get(k) for k in head])
    ws.freeze_panes = 'A2'
    for col, width in zip('ABCDE', (38, 16, 20, 13, 26)):
        ws.column_dimensions[col].width = width
    wb.save(a.file)
    live = sum(1 for c in rows if not c.get('unsubscribed'))
    print(f'{len(rows)} contacts ({live} subscribed, {len(rows) - live} unsubscribed) -> {a.file}')


def cmd_contacts_template(a):
    path = Path(a.file)
    if path.exists():
        sys.exit(f'{path} already exists; not overwriting it.')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = 'Contacts'
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0E1B2A')
    for col, width in zip('ABCDEFGH', (36, 16, 20, 30, 11, 22, 14, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    for rng, options in (('E2:E5000', '"en,fr"'), ('F2:F5000', '"' + ','.join(SOURCES) + '"'), ('G2:G5000', '"unsubscribe"')):
        dv = DataValidation(type='list', formula1=options, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)
    g = wb.create_sheet('How to use')
    lines = [
        ('Heirstone newsletter contacts', True),
        ('', False),
        ('One row per person. Only email and source are required.', False),
        ('email: their work email. first_name: used in the greeting ("Dear Amina"). organisation: kept in Resend for reference.', False),
        ('language: en or fr. Sets which edition they receive (blank = en).', False),
        ('source: why we may email them. Rows without one are skipped:', False),
    ] + [(f'    {k}: {v}', False) for k, v in SOURCES.items()] + [
        ('action: leave blank to add or update; write unsubscribe to stop sending to someone (they stay on file so they are never re-added).', False),
        ('notes: for you only; not uploaded.', False),
        ('', False),
        ('Do not add people who have not agreed to hear from us (BD target lists, RocketReach or LinkedIn exports, bought lists).', True),
        ('Write to them one to one and invite them to download a report instead: that signs them up properly.', False),
        ('Bounces and spam complaints from cold lists can get the Resend account suspended, and that account also sends', False),
        ('the report emails and contact-form messages for Heirstone and five other ventures.', False),
        ('', False),
        ('To upload: save this file, then run  python newsletter/nl.py contacts sync "<path to this file>"', False),
        ('(or ask Claude to "sync the newsletter contacts"). Rows can stay in the file: re-running is safe.', False),
        ('Website sign-ups go straight into Resend and do not appear here. To see the full live list:', False),
        ('python newsletter/nl.py contacts export "<path>.xlsx"', False),
    ]
    for text, bold in lines:
        g.append([text])
        if bold:
            g.cell(g.max_row, 1).font = Font(bold=True)
    g.column_dimensions['A'].width = 130
    for r in g.iter_rows():
        r[0].alignment = Alignment(wrap_text=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f'Created {path}')


# ---------------------------------------------------------------- command: setup
PROPERTIES = ['organisation', 'language', 'site_language', 'source', 'first_report', 'last_report', 'signed_up_at', 'last_report_at']


def cmd_setup(a):
    cfg = load_config()
    ids = cfg.setdefault('resend', {})
    ids.setdefault('topics', {})

    def listing(path):
        s, res = resend('GET', path)
        if s >= 300:
            sys.exit(f'Resend refused {path} ({s}): {res}. The API key needs full access.')
        return res.get('data', [])

    if not ids.get('segment_id'):
        seg = next((x for x in listing('/segments') if x.get('name') == a.segment), None)
        if not seg:
            s, seg = resend('POST', '/segments', {'name': a.segment})
            print(f'segment "{a.segment}": created')
        ids['segment_id'] = seg['id']
    existing = {x.get('name'): x for x in listing('/topics')}
    for lang, (tname, default, desc) in {
        'en': (f"{cfg['name']['en']} (English)", 'opt_in', 'Heirstone Consulting research newsletter, English edition'),
        'fr': (f"{cfg['name']['fr']} (français)", 'opt_out', 'Lettre de recherche de Heirstone Consulting, édition française'),
    }.items():
        if ids['topics'].get(lang):
            continue
        t = existing.get(tname)
        if not t:
            s, t = resend('POST', '/topics', {'name': tname, 'default_subscription': default, 'description': desc, 'visibility': 'private'})
            if s >= 300:
                sys.exit(f'Could not create topic {tname}: {t}')
            print(f'topic "{tname}": created ({default})')
        ids['topics'][lang] = t['id']
    have = {x.get('key') for x in listing('/contact-properties')}
    for key in PROPERTIES:
        if key not in have:
            s, res = resend('POST', '/contact-properties', {'key': key, 'type': 'string'})
            print(f'field {key}: {"created" if s < 300 else "FAILED " + str(res)}')
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(f'IDs saved in {CONFIG_FILE.relative_to(ROOT)}. Commit and deploy it to switch on automatic sign-ups.')


def main():
    for stream in (sys.stdout, sys.stderr):  # Windows consoles default to a legacy code page
        stream.reconfigure(encoding='utf-8', errors='replace')
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    s = sub.add_parser('render'); s.add_argument('issue'); s.set_defaults(fn=cmd_render)
    s = sub.add_parser('test'); s.add_argument('issue'); s.add_argument('--to', required=True)
    s.add_argument('--name', default='', help='first name to show in the greeting'); s.add_argument('--lang', help='en, fr or en,fr'); s.set_defaults(fn=cmd_test)
    s = sub.add_parser('draft'); s.add_argument('issue'); s.add_argument('--force', action='store_true')
    s.add_argument('--skip-image-check', action='store_true'); s.set_defaults(fn=cmd_draft)
    s = sub.add_parser('setup'); s.add_argument('--segment', default='Heirstone Newsletter'); s.set_defaults(fn=cmd_setup)
    c = sub.add_parser('contacts').add_subparsers(dest='sub', required=True)
    s = c.add_parser('sync'); s.add_argument('file'); s.add_argument('--dry-run', action='store_true'); s.set_defaults(fn=cmd_contacts_sync)
    s = c.add_parser('export'); s.add_argument('file'); s.set_defaults(fn=cmd_contacts_export)
    s = c.add_parser('template'); s.add_argument('file'); s.set_defaults(fn=cmd_contacts_template)
    a = ap.parse_args()
    a.fn(a)


if __name__ == '__main__':
    main()
